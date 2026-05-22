"""
finRAG API — Ingest 路由
========================
POST /api/v1/ingest — 上传 PDF，触发解析+索引
GET /api/v1/ingest/{task_id} — 查询入库任务状态
"""
import json
import logging
import time
import uuid
from pathlib import Path

from fastapi import APIRouter, File, HTTPException, UploadFile

from backend.config import settings
from backend.models.schemas import IngestResponse, IngestStatus

logger = logging.getLogger("api.ingest")
router = APIRouter(prefix="/api/v1", tags=["Ingest"])

# 内存任务状态（生产环境应换 Redis）
_tasks: dict[str, IngestStatus] = {}


@router.post("/ingest", response_model=IngestResponse)
async def ingest_pdf(
    file: UploadFile = File(...),
    collections: str = "pro",
    generate_outline: bool = False,
):
    """
    上传文档并触发解析+索引流水线。
    支持格式：.pdf .docx .pptx .xlsx

    参数:
    - collections: 目标集合，逗号分隔。例如 "pro"（联网库）或 "safe"（安全库）或 "pro,safe"（双库）
    - generate_outline: 是否生成结构纲要（会消耗 DeepSeek API token）
    
    处理流程（后台异步）：
    1. 保存上传文件到 data/raw/
    2. 解析为 Markdown
    3. 调用 ingestion_v2.py Markdown → 目标 Qdrant 集合
    """
    ALLOWED_EXTS = {".pdf", ".docx", ".pptx", ".xlsx", ".md", ".txt"}
    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名为空")
    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise HTTPException(status_code=400, detail=f"不支持的文件格式: {ext}，仅支持 {', '.join(sorted(ALLOWED_EXTS))}")

    task_id = str(uuid.uuid4())[:8]
    # 安全净化：只取 basename，拒绝路径遍历
    safe_name = Path(file.filename).name
    if safe_name != file.filename or ".." in safe_name or "/" in safe_name or "\\" in safe_name:
        raise HTTPException(status_code=400, detail="文件名不合法")
    save_path = settings.RAW_DIR / safe_name

    # 检查重复
    if save_path.exists():
        return IngestResponse(
            status="error",
            task_id=task_id,
            message=f"文件已存在: {file.filename}",
            file_name=file.filename,
            error="文件已存在，请先删除旧文件或重命名",
        )

    # 保存文件
    try:
        settings.RAW_DIR.mkdir(parents=True, exist_ok=True)
        content = await file.read()
        save_path.write_bytes(content)
        # 防御性验证：确认文件确实在 RAW_DIR 下
        resolved = save_path.resolve()
        raw_resolved = settings.RAW_DIR.resolve()
        if not str(resolved).startswith(str(raw_resolved)):
            save_path.unlink(missing_ok=True)
            raise HTTPException(status_code=400, detail="文件路径异常")
    except Exception as e:
        return IngestResponse(
            status="error",
            task_id=task_id,
            message=f"文件保存失败: {e}",
            file_name=file.filename,
            error=str(e),
        )

    # 初始化任务状态
    _tasks[task_id] = IngestStatus(
        task_id=task_id,
        status="queued",
        file_name=file.filename,
        message="文件已保存，解析队列中",
    )

    # 后台异步处理（生产环境使用 Celery）
    import asyncio
    asyncio.create_task(_process_ingest(task_id, save_path, collections, generate_outline))

    return IngestResponse(
        status="processing",
        task_id=task_id,
        message=f"文件 {file.filename} 已加入解析队列",
        file_name=file.filename,
    )


def _parse_office_document(path: Path) -> str:
    """将 Office 文档（.docx .pptx .xlsx）解析为 Markdown 文本。"""
    ext = path.suffix.lower()
    parts = []

    if ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            style = p.style.name.lower() if p.style else ""
            if "heading" in style or "title" in style:
                level = min(int(c) for c in style if c.isdigit()) if any(c.isdigit() for c in style) else 1
                parts.append(f"{'#' * min(level, 6)} {t}")
            else:
                parts.append(t)
        for table in doc.tables:
            parts.append("")
            for i, row in enumerate(table.rows):
                cells = " | ".join(c.text.strip() for c in row.cells)
                parts.append(f"| {cells} |")
                if i == 0:
                    parts.append("| " + " | ".join("---" for _ in row.cells) + " |")
            parts.append("")

    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        for slide in prs.slides:
            parts.append(f"---\n")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        t = p.text.strip()
                        if t:
                            parts.append(t)
                if shape.has_table:
                    table = shape.table
                    for i, row in enumerate(table.rows):
                        cells = " | ".join(c.text.strip() for c in row.cells)
                        parts.append(f"| {cells} |")
                        if i == 0:
                            parts.append("| " + " | ".join("---" for _ in row.cells) + " |")
                    parts.append("")

    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"\n## {sheet_name}\n")
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            # Skip empty sheets
            rows = [r for r in rows if any(c is not None for c in r)]
            if not rows:
                continue
            for i, row in enumerate(rows):
                cells = " | ".join(str(c) if c is not None else "" for c in row)
                parts.append(f"| {cells} |")
                if i == 0:
                    parts.append("| " + " | ".join("---" for _ in row) + " |")
            parts.append("")
        wb.close()

    return "\n".join(parts)


async def _process_ingest(task_id: str, file_path: Path, collections: str = "pro", generate_outline: bool = False):
    """后台执行解析+索引（异步）。支持多集合索引。"""
    from backend.models.schemas import IngestStage

    ext = file_path.suffix.lower()
    target_modes = [m.strip() for m in collections.split(",") if m.strip() in ("pro", "safe")]
    try:
        _tasks[task_id].status = "processing"
        _tasks[task_id].stage = IngestStage.PARSING

        import sys
        sys.path.insert(0, str(settings.BASE_DIR))

        # ═══════════════ Phase 1: 文档 → Markdown ═══════════════
        _tasks[task_id].progress = 0.1
        _tasks[task_id].message = "正在解析文档..."

        if ext == ".pdf":
            _tasks[task_id].message = "正在解析 PDF（第 1 页）..."
            from scripts.ingestion import PDFParser, MarkdownCleaner
            from scripts.ingestion import Config as IngestionConfig
            icfg = IngestionConfig()
            icfg.RAW_DIR = settings.RAW_DIR
            icfg.PARSED_DIR = settings.PARSED_DIR
            parser = PDFParser(pdf_type="text")
            cleaner = MarkdownCleaner(icfg)
            md_text, error, page_count, status = parser.parse(str(file_path))
            if status == "error":
                raise RuntimeError(f"PDF 解析失败: {error}")
            md_text = cleaner.clean(md_text)
            _tasks[task_id].message = f"PDF 解析完成（共 {page_count} 页），正在清理格式..."
        elif ext in (".md", ".txt"):
            _tasks[task_id].message = f"正在读取 {ext.upper()}..."
            md_text = file_path.read_text(encoding="utf-8")
            page_count = md_text.count("\n\n") + 1
        else:
            _tasks[task_id].message = f"正在解析 {ext.upper()}..."
            md_text = _parse_office_document(file_path)
            page_count = md_text.count("\n---\n") + 1

        md_path = settings.PARSED_DIR / f"{file_path.stem}.md"
        md_path.write_text(md_text, encoding="utf-8")
        char_count = len(md_text)
        logger.info(f"  ✅ {ext.upper()} → Markdown: {md_path.name} ({char_count} 字)")

        _tasks[task_id].progress = 0.4
        _tasks[task_id].message = f"解析完成（{char_count} 字），正在索引向量库..."
        _tasks[task_id].stage = IngestStage.INDEXING

        # ═══════════════ Phase 2: Markdown → Qdrant ═══════════════
        from scripts.ingestion_v2 import (
            Config as V2Config,
            MetadataExtractor,
            ParentChildSplitter,
            EmbeddingEngine,
            QdrantIndexer,
        )

        v2cfg = V2Config()
        v2cfg.PARSED_DIR = settings.PARSED_DIR
        v2cfg.QDRANT_DB_PATH = settings.QDRANT_PATH
        # 尊重后端配置：远程 Qdrant 时用 remote 模式
        qdrant_host = settings.QDRANT_HOST
        if qdrant_host and qdrant_host not in ("localhost", "127.0.0.1"):
            v2cfg.QDRANT_MODE = "remote"
            v2cfg.QDRANT_HOST = qdrant_host
            v2cfg.QDRANT_PORT = settings.QDRANT_PORT
            logger.info(f"  🔌 使用远程 Qdrant: {qdrant_host}:{settings.QDRANT_PORT}")
        else:
            v2cfg.QDRANT_MODE = "disk"
            logger.info(f"  🔌 使用本地 Qdrant: {settings.QDRANT_PATH}")

        _tasks[task_id].progress = 0.45
        _tasks[task_id].message = "正在提取元数据..."

        meta_extractor = MetadataExtractor(v2cfg)
        metadata = meta_extractor.extract(md_path, md_text)
        logger.info(f"      🏷️  元数据: {metadata['company']} | {metadata['year']} | {metadata['doc_type']}")

        _tasks[task_id].progress = 0.48
        _tasks[task_id].message = "正在分块..."

        splitter = ParentChildSplitter(v2cfg)
        parents = splitter.split(md_text)
        children_count = sum(len(p.children) for p in parents)
        logger.info(f"      ✂️  {len(parents)} 父块, {children_count} 子块")

        if not parents:
            raise RuntimeError("文档分块后无有效内容")

        _tasks[task_id].progress = 0.50
        _tasks[task_id].message = "正在加载嵌入模型..."

        embedder = EmbeddingEngine(v2cfg)
        embedder._load()

        # 对每个目标集合执行索引
        collection_name_map = {
            "pro": settings.QDRANT_COLLECTION_PRO,
            "safe": settings.QDRANT_COLLECTION_SAFE,
        }
        union_task_msg = ""
        n_parents = 0
        modes_done = 0
        # 复用后端 retriever 的 Qdrant 连接，避免 disk 锁冲突
        existing_client = None
        try:
            from backend.services.rag_engine import _get_retriever
            r = _get_retriever(target_modes[0] if target_modes else "pro")
            if r._qdrant_client is not None:
                existing_client = r._qdrant_client
                logger.info("  🔗 复用后端 retriever 的 Qdrant 连接")
        except Exception as e:
            logger.warning(f"  ⚠️  无法获取现有 Qdrant 连接: {e}")

        for target_mode in target_modes:
            target_col = collection_name_map.get(target_mode, settings.QDRANT_COLLECTION_PRO)
            _tasks[task_id].message = f"正在编码并写入 {target_mode} 集合..."

            v2cfg.QDRANT_COLLECTION = target_col
            indexer = QdrantIndexer(v2cfg, existing_client=existing_client)
            indexer._connect()
            indexer.ensure_collection()
            n_parents = indexer.index_document(metadata, parents, embedder)
            indexer.close()
            logger.info(f"  ✅ → {target_col}: {n_parents} 父块")
            union_task_msg += f"{target_col}: {n_parents}父块 + {children_count}子块; "
            modes_done += 1
            # 进度从 0.55 均匀推进到 0.75（按集合数分摊）
            _tasks[task_id].progress = 0.55 + (modes_done / len(target_modes)) * 0.20

        _tasks[task_id].progress = 0.75
        _tasks[task_id].message = "向量索引完成，正在刷新文档列表..."
        _tasks[task_id].stage = IngestStage.FINALIZING

        # ═══════════════ Phase 3: 后处理 ═══════════════
        _tasks[task_id].progress = 0.80
        _tasks[task_id].message = "正在刷新文档列表缓存..."

        try:
            from backend.routers.sources import refresh_sources_cache
            refresh_sources_cache()
            logger.info("  ✅ sources_index.json 已刷新")
        except Exception as e:
            logger.warning(f"  ⚠️  刷新 sources 缓存失败: {e}")

        _tasks[task_id].progress = 0.88
        _tasks[task_id].message = "正在重建全文索引..."

        try:
            from scripts.retriever import HybridRetriever
            retriever = HybridRetriever(v2cfg)
            retriever.build_bm25_index()
            logger.info("  ✅ BM25 索引已重建")
        except Exception as e:
            logger.warning(f"  ⚠️  BM25 重建失败: {e}")

        # 清空查询缓存
        _tasks[task_id].progress = 0.95
        _tasks[task_id].message = "正在清空查询缓存..."
        try:
            from backend.services.cache import query_cache
            query_cache.invalidate("q:")
            logger.info("  🧹 查询缓存已清空")
        except Exception:
            pass

        _tasks[task_id].progress = 1.0
        _tasks[task_id].stage = IngestStage.DONE
        _tasks[task_id].status = "completed"
        summary = union_task_msg.strip("; ") if union_task_msg else f"{n_parents} 父块 + {children_count} 子块"
        _tasks[task_id].message = f"文档已就绪：{summary}"
        _tasks[task_id].chunks_count = n_parents + children_count

        # ═══════════════ Phase 4: 结构纲要（可选的离线任务）═══════════
        if generate_outline:
            _tasks[task_id].message = f"文档已就绪，正在离线生成结构纲要..."
            try:
                # 对每个目标模式分别生成纲要
                for target_mode in target_modes:
                    import subprocess
                    script_path = settings.BASE_DIR / "scripts" / "generate_outlines.py"
                    logger.info(f"  📋 开始生成 {target_mode} 模式大纲...")
                    result = subprocess.run(
                        [sys.executable, str(script_path), "--mode", target_mode, "--doc", file_path.stem],
                        capture_output=True, text=True, timeout=600,
                    )
                    if result.returncode != 0:
                        logger.warning(f"  ⚠️  纲要生成失败 ({target_mode}): {result.stderr[:200]}")
                    else:
                        logger.info(f"  ✅ 纲要生成完成 ({target_mode})")
                _tasks[task_id].message = f"文档已就绪（含结构纲要）"
            except Exception as e:
                logger.warning(f"  ⚠️  纲要生成异常: {e}")
                _tasks[task_id].message = f"文档已就绪（纲要生成失败: {e}）"

    except Exception as e:
        logger.error(f"Ingest task {task_id} failed: {e}", exc_info=True)
        _tasks[task_id].status = "error"
        _tasks[task_id].message = f"处理失败: {e}"
        _tasks[task_id].error = str(e)


@router.get("/ingest/{task_id}", response_model=IngestStatus)
async def get_ingest_status(task_id: str):
    """查询入库任务状态。"""
    task = _tasks.get(task_id)
    if not task:
        raise HTTPException(status_code=404, detail=f"任务 {task_id} 不存在")
    return task
